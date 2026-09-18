from globals import *
from my_functions import *
import my_functions

import logging
import os
from datetime import datetime
import sys
import signal

# ============================================================
# BASIC LOGGING CONFIGURATION
# ============================================================

LOG_DIR = Path(__file__).resolve().parent / "logs"
LOG_DIR.mkdir(parents=True, exist_ok=True)

# Nombre del script
script_name = os.path.splitext(os.path.basename(__file__))[0]

LOG_FILE = LOG_DIR / (
    f"{script_name}_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.log"
)

# Forzar creación inmediata del archivo
LOG_FILE.touch()  # ← agrega esta línea

# Handler con flush inmediato
_file_handler = logging.FileHandler(LOG_FILE, mode="w", encoding="utf-8")
_file_handler.setFormatter(logging.Formatter(
    "%(asctime)s | %(levelname)s | %(filename)s:%(lineno)d | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S"
))

# Subclase simple que hace flush después de cada línea
class _ImmediateFileHandler(logging.FileHandler):
    def emit(self, record):
        super().emit(record)
        self.stream.flush()

logging.basicConfig(
    force=True,
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(filename)s:%(lineno)d | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    handlers=[
        _ImmediateFileHandler(LOG_FILE, mode="w", encoding="utf-8"),
        logging.StreamHandler(sys.__stdout__),
    ],
)

logger = logging.getLogger(__name__)
logger.info("LOG FILE PATH: %s", LOG_FILE)
logger.info("Logger initialized")

class StreamToLogger:
    def __init__(self, logger, level, original_stream):
        self.logger = logger
        self.level = level
        self.original_stream = original_stream
    def write(self, message):
        if message.strip():
            self.logger.log(self.level, message.strip())
            self.original_stream.write(message)
            self.original_stream.flush()
    def flush(self):
        self.original_stream.flush()

sys.stdout = StreamToLogger(logger, logging.INFO, sys.__stdout__)
sys.stderr = StreamToLogger(logger, logging.ERROR, sys.__stderr__)

# visual comment: capture any uncaught exception and log it
def global_exception_handler(exc_type, exc_value, exc_traceback):
    logger.exception(
        "Uncaught exception",
        exc_info=(exc_type, exc_value, exc_traceback),
    )

sys.excepthook = global_exception_handler

logger.info("TEST LOG INITIALIZATION")
logger.info("Logger test executed")

# visual comment: handle Ctrl+C and forced exits
def shutdown_handler(signum, frame):
    logger.info("Shutdown signal received. Stopping OBS recording...")
    try:
        end_record_obs()
    except Exception as e:
        logger.error(f"Error stopping OBS: {e}")
    finally:
        sys.exit(0)

signal.signal(signal.SIGINT, shutdown_handler)   # Ctrl+C
signal.signal(signal.SIGTERM, shutdown_handler)  # kill / cierre externo


#import atexit

#atexit.register(end_record_obs)

from openpyxl import Workbook, load_workbook

silver_conquest = True

try:
    subprocess.run(
        ["taskkill", "/im", "snap.exe", "/f"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    logger.info("snap.exe terminated if it was running")
    

    resize_window_by_tuple(
        title_contains=["exp_farm_conquest"],  # nombre de tu terminal/script
        region=globals.REGION_TOTAL_MINUS_MS_WINDOW,
    )

    start_obs_then_record()

    os.system('start steam://rungameid/1997040')
    logger.info("Steam launch command sent")

    resize_window_by_tuple(['SNAP'], globals.REGION_MS_WINDOW, delay=10)    
    
    #######################################################################################################
    my_functions.login('john')
    
    search_initial_imgs()
    
    select_deck("3turns")
    
    click_any_MS_img(r'images/start_screen/maybe_later.PNG', timeout=2)
    click_any_MS_img(r'images/start_screen/maybe_later2.PNG', timeout=0.01)
        
    print("Abriendo modo de juego Conquest...")
    
    click_any_MS_img(r'images/unselected/game_modes_unselected.PNG', timeout=3)
    click_any_MS_img(r'images/games_mode_selected.PNG', timeout=3)
    click_any_MS_img(r'images/stages/conquest/conquest.PNG', timeout=3)

    pg.moveTo(400,400)

    time.sleep(3)
    
    if silver_conquest:
        pg.scroll(-100)
    else:
        print("Desplazando hacia abajo hasta llegar a Proving Grounds...")
        pg.scroll(-2000)

    logger.info("Entrando al bucle de juego...")

    game = 0
    previous_xp = get_postgame_xp_value()

    while True:
    
        start_time = time.time()  # visual comment: start time per game
        
        if click_any_MS_img(r'images/stages/conquest/enter.PNG', timeout=10):
            pg.move(0, -500)
        else:
            presionar_esc()
            click_any_MS_img(r'images/stages/conquest/conquest.PNG', timeout=3)
            if click_any_MS_img('images/next_claim_ticket_3.PNG', timeout=15):
                click_any_MS_img('images/claim!_after_gain_ticket_conquest_2.PNG', delay_click=.6)
            
            click_any_MS_img(r'images/stages/conquest/enter.PNG', timeout=10)
            pg.move(0, -500)

        time.sleep(5)

        click_any_MS_img(r'images/stages/conquest/play_conquest.PNG')
        click_any_MS_img(r'images/stages/conquest/confirm_conquest.PNG')
        
        
        if search_MS_img('images/matchmaking_has_been_interrupted.png'):
            click_any_MS_img('images/red_x_btn.png')
            click_any_MS_img(r'images/stages/conquest/play_conquest.PNG')
            click_any_MS_img(r'images/stages/conquest/confirm_conquest.PNG')
        

        logger.info("Entrando al bucle de juego...")
        
        previous_time = 0

        while True:

            if search_MS_img('images/next_after_concede_conquest_2.PNG', verbose=False):
                logger.info("Fin de la partida. Mostrando resultados.")
                click_any_MS_img('images/next_after_concede_conquest_2.PNG')

                logger.debug("Verificando si fue victoria...")

                if search_MS_img('images/battle_victory_conquest.PNG', timeout=3, verbose=False):
                    logger.info("¡Victoria detectada! Reclamando recompensas...")
                    click_any_MS_img('images/next_conquest.PNG')
                    time.sleep(0.5)
                    click_any_MS_img('images/next_conquest.PNG')
                    click_any_MS_img('images/next_claim_ticket_3.PNG', timeout=25)
                    click_any_MS_img('images/claim!_after_gain_ticket_conquest_2.PNG', delay_click=.6)
                    pg.move(0, -500)
                    time.sleep(1)
                    pg.scroll(-1000)
                    break

                else:
                    logger.info("Derrota detectada. No se reclama nada.")
                    click_any_MS_img('images/next_conquest.PNG')
                    time.sleep(0.5)
                    click_any_MS_img('images/next_conquest.PNG')
                    break

            elif search_MS_img('images/battle_ready.PNG', verbose=False):
                logger.info("Fin de ronda. Hay más rondas por jugar.")
                click_any_MS_img('images/battle_ready.PNG')
                time.sleep(3)
                continue

            elif end_turn := search_MS_img('images/stages/play_stage/end_turn.PNG', verbose=False):
                logger.debug("Tu turno. Jugando cartas...")
                play_cards_ahk(globals.POSITIONS_MOUSE_3_TURNS)
                pg.click(end_turn)

                if search_MS_img('images/stages/play_stage/flechas.png', verbose=False):
                    presionar_esc()

                pg.click(end_turn)

            elif search_MS_img('images/stages/play_stage/playing3_6_2.png', confidence=0.95, verbose=False):
                logger.info("Turno 3+ detectado. Retirándose para ganar EXP.")
                presionar_esc()
                click_any_MS_img('images/stages/play_stage/retreat_now.PNG', delay_click=.1)
                time.sleep(5)

            if search_MS_img('images/you_have_encountered_an_error.png', verbose=False):
                logger.warning("Error del juego detectado. Reiniciando...")
                pg.hotkey('alt', 'f4')
                os.system('start steam://rungameid/1997040')
                resize_window_by_tuple('SNAP', globals.REGION_MS_WINDOW, delay=10)

                logger.info("Abriendo modo de juego Conquest...")
                click_any_MS_img(r'images/unselected/game_modes_unselected.PNG')
                click_any_MS_img(r'images/stages/conquest/conquest.PNG', timeout=3)

                pg.moveTo(400,400)

                logger.info("Desplazando hacia abajo para ver entradas disponibles...")
                time.sleep(3)
                pg.scroll(-2000)

                logger.info("Entrando a la partida...")

            if search_MS_img('images/aw_snap_only.png', verbose=False):
                logger.warning("AwSnap detectado.")
                click_any_MS_img('images/red_x_btn')
                break

            if search_MS_img('images/battle_victory_conquest.PNG', verbose=False):
                logger.info("¡Victoria detectada! Reclamando recompensas...")
                click_any_MS_img('images/next_conquest.PNG')
                time.sleep(0.5)
                click_any_MS_img('images/next_conquest.PNG')
                click_any_MS_img('images/next_claim_ticket_3.PNG', timeout=25)
                click_any_MS_img('images/claim!_after_gain_ticket_conquest_2.PNG', delay_click=.6)
                pg.move(0, -500)
                time.sleep(1)
                pg.scroll(-1000)
                break

            time.sleep(0.5)
    
        print(f'Game #{game+1} finalizada\n')

        end_time = time.time()
        minutes_game = (end_time - start_time) / 60
        
        print(f"Tiempo total de la partida: {minutes_game:.1f} min")

 
        xp_value = get_postgame_xp_value()

        xp_gained = xp_value - previous_xp
        previous_xp = xp_value

        print(f"XP obtenida en este game: {xp_gained}")
            
            
        
        # Actualizar tabla con exp del game ultimo        
        # Usar la variable global definida dentro del módulo
        
        excel_path = fr"Z:\SHARED FILES VM\conquest_xp_{my_functions.CURRENT_ACCOUNT}.xlsx"
        print(my_functions.CURRENT_ACCOUNT)
        
        today_sheet = datetime.now().strftime("%Y-%m-%d")

        if not os.path.exists(excel_path):
            wb = Workbook()
            ws = wb.active
            ws.title = today_sheet
            ws.append(["Game", "XP", "Tiempo en minutos", "Hora fin"])
        else:
            wb = load_workbook(excel_path)

            if today_sheet in wb.sheetnames:
                ws = wb[today_sheet]
            else:
                ws = wb.create_sheet(title=today_sheet)
                ws.append(["Game", "XP", "Tiempo en minutos", "Hora fin"])

        # Número de game basado en filas
        game_number = ws.max_row
        
        end_hour = datetime.now().strftime("%H:%M")
    
        # Agregar nueva fila
        ws.append([game_number, xp_gained, round(minutes_game), end_hour])       

        # visual comment: calculate totals
        total_xp = 0
        total_time = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1]:
                total_xp += row[1]
            if row[2]:
                total_time += row[2]

        # visual comment: write totals in first row (side columns)
        ws["F1"] = "TOTAL XP"
        ws["G1"] = total_xp

        ws["H1"] = "TOTAL TIME"
        ws["I1"] = round(total_time, 1)
        
        # Guardar
        wb.save(excel_path)
        
        game += 1
        
        if xp_value >= 3700:
            print('Se alcanzo 3700 o más de xp')
            beep_forever()


except Exception as e:
    logger.exception("Unhandled exception occurred")

    input("\nPresiona ENTER para cerrar...")
    
finally:
    logger.info("Finalizing script. Stopping OBS recording...")
    try:
        end_record_obs()
    except Exception as e:
        logger.error(f"Error stopping OBS in finally: {e}")